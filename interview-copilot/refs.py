"""Reference library: a small local search index over your study notes.

Indexes .md, .txt, .pdf and .docx files under a folder, splits them into sections,
and finds the sections that best match an interview question (BM25 ranking).
Everything runs locally; only the few matching sections are sent with a question.
"""

import io
import math
import os
import re
import threading
import time
from collections import Counter, defaultdict
from pathlib import Path

DOC_TYPES = (".md", ".txt", ".pdf", ".docx")
# Folders never indexed: hidden folders (.git, .trunk ...) are skipped too.
DEFAULT_SKIP = {"node_modules", ".venv", "venv", "__pycache__", "coverage", "interview-copilot", "Resume"}
MAX_FILE_BYTES = 15 * 1024 * 1024
CHUNK_CHARS = 1400

STOPWORDS = set("""a an and are as at be been but by can could do does did for from had has have how i if in
into is it its me my of on or our so such than that the their them then there these they this to us
was we were what when where which while who why will with would you your about also any each just
more most other some very via vs using use used tell explain describe walk through give example
know like let lets please talk implement design handle approach ensure make work difference
between what's whats you'd would should""".split())

# Words that never justify a search on their own ("Why is that better?" is a follow-up, not a topic)
GENERIC = set("""better best good bad worse example benefit advantage disadvantage
alternative challenge problem issue scenario experience important mean meaning reason detail more
else other thing way option case question answer point idea opinion choice""".split())

# Interview shorthand -> words your notes are likely to use
SYNONYMS = {
    "k8s": ["kubernetes"], "kubernetes": ["k8s"], "eks": ["kubernetes"], "aks": ["kubernetes"],
    "gke": ["kubernetes"], "tf": ["terraform"], "iac": ["terraform", "infrastructure"],
    "cicd": ["pipeline"], "ci": ["pipeline"], "cd": ["pipeline", "deployment"],
    "lb": ["load", "balancer"], "alb": ["load", "balancer"], "nlb": ["load", "balancer"],
    "sg": ["security", "group"], "vpc": ["network"], "rds": ["database"], "db": ["database"],
    "iam": ["identity", "access"], "rbac": ["role", "access"], "dr": ["disaster", "recovery"],
    "ha": ["availability"], "s3": ["bucket", "storage"], "kms": ["encryption", "key"],
    "ec2": ["instance"], "llm": ["model"], "rag": ["retrieval"], "sre": ["reliability"],
    "finops": ["cost"], "devsecops": ["security", "pipeline"], "helm": ["chart"],
}

TOKEN_RE = re.compile(r"[a-z0-9]+(?:[.+#][a-z0-9]+)*")


def tokenize(text):
    out = []
    for t in TOKEN_RE.findall(text.lower()):
        if t in STOPWORDS or (len(t) < 2 and not t.isdigit()):
            continue
        if len(t) > 4 and t.endswith("s") and not t.endswith("ss"):
            t = t[:-1]  # buckets -> bucket
        out.append(t)
    return out


def extract_text(filename, data):
    """Plain text from a .docx, .pdf, .txt or .md file."""
    name = filename.lower()
    if name.endswith(".docx"):
        import docx
        doc = docx.Document(io.BytesIO(data))
        lines = []
        for p in doc.paragraphs:
            style = (p.style.name or "").lower() if p.style is not None else ""
            text = p.text
            if text.strip() and style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "2"
                text = "#" * min(int(level), 6) + " " + text  # keep headings for sectioning
            lines.append(text)
        for table in doc.tables:  # many templates put content in tables
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    if cell.text.strip() and cell.text not in cells:
                        cells.append(cell.text)
                lines.append(" | ".join(cells))
        text = "\n".join(lines)
    elif name.endswith(".pdf"):
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError("Please upload a .docx, .pdf or .txt file.")
    out, blank = [], False
    for line in text.splitlines():
        line = line.rstrip()
        if not line.strip():
            if not blank:
                out.append("")
            blank = True
        else:
            out.append(line)
            blank = False
    text = "\n".join(out).strip()
    if not text:
        raise ValueError("No text found in that file. If it's a scanned PDF, upload the .docx instead.")
    return text


def split_sections(text):
    """Split into ~CHUNK_CHARS sections, remembering the heading path of each."""
    heads, chunks, buf, buf_head = [], [], [], ""

    def flush():
        body = "\n".join(buf).strip()
        if len(body) > 40:
            chunks.append((buf_head, body))
        buf.clear()

    in_code = False
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
        m = None if in_code else re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            level, title = len(m.group(1)), m.group(2).strip().strip("#").strip()
            heads[:] = [h for h in heads if h[0] < level] + [(level, title)]
            if sum(len(x) for x in buf) > 250:
                flush()
                buf_head = " › ".join(h[1] for h in heads)
            elif not buf:
                buf_head = " › ".join(h[1] for h in heads)
            buf.append(line)
            continue
        buf.append(line)
        if sum(len(x) + 1 for x in buf) >= CHUNK_CHARS and not line.strip() and not in_code:
            flush()
            buf_head = " › ".join(h[1] for h in heads)
    flush()
    # very long sections without blank lines: hard split
    out = []
    for head, body in chunks:
        while len(body) > CHUNK_CHARS * 2:
            cut = body.rfind("\n", 0, CHUNK_CHARS * 2)
            cut = cut if cut > CHUNK_CHARS // 2 else CHUNK_CHARS * 2
            out.append((head, body[:cut].strip()))
            body = body[cut:].strip()
        out.append((head, body))
    return out


def section_label(chunk):
    """Short readable name for a section: its own heading, or its first line if headings are generic."""
    def clean(s):
        s = re.sub(r"<[^>]+>", "", s)                    # html tags
        s = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", s)   # [text](link) -> text
        return s.replace("**", "").strip().lstrip("#").strip().strip("*").strip()

    last = clean(chunk["section"].split(" › ")[-1]) if chunk["section"] else ""
    first = ""
    for line in chunk["text"].splitlines():
        if line.startswith(("    ", "\t")) or any(ch in line for ch in "{}=;") or line.strip().startswith(("```", "$", "|")):
            continue  # code or table rows make poor labels
        line = clean(line)
        if len(line) > 8:
            first = line
            break
    label = last if last and len(last) > 12 and not first.lower().startswith(last.lower()[:20]) else first or last
    return label[:90] + ("…" if len(label) > 90 else "")


class Library:
    def __init__(self, root, skip=DEFAULT_SKIP):
        self.root = Path(root)
        self.skip = set(skip)
        self.lock = threading.Lock()
        self.chunks = []          # [{file, section, text, len}]
        self.postings = {}        # term -> [(chunk_index, tf)]
        self.idf = {}
        self.avg_len = 1.0
        self.topics = set()
        self.files = 0
        self.errors = []
        self.built_at = None
        self.build_seconds = None
        self.building = False

    # ---------------------------------------------------------------- build
    def _doc_paths(self):
        if not self.root.is_dir():
            return []
        paths = []
        for dirpath, dirnames, filenames in os.walk(self.root):
            dirnames[:] = sorted(d for d in dirnames if not d.startswith(".") and d not in self.skip)
            for f in sorted(filenames):
                if f.lower().endswith(DOC_TYPES) and not f.startswith("~$"):
                    paths.append(Path(dirpath) / f)
        return paths

    def build(self):
        with self.lock:
            if self.building:
                return
            self.building = True
        started = time.time()
        chunks, errors, files = [], [], 0
        for path in self._doc_paths():
            try:
                if path.stat().st_size > MAX_FILE_BYTES:
                    continue
                text = extract_text(path.name, path.read_bytes())
            except Exception as e:  # unreadable or empty file: note it and move on
                errors.append(f"{path.relative_to(self.root)}: {e}")
                continue
            files += 1
            rel = str(path.relative_to(self.root)).replace("\\", "/")
            title_words = rel.rsplit(".", 1)[0].replace("/", " ").replace("_", " ").replace("-", " ")
            for head, body in split_sections(text):
                lines = [ln for ln in body.splitlines() if ln.strip()]
                if lines and sum("](#" in ln for ln in lines) / len(lines) > 0.4:
                    continue  # a table of contents: just links, no content
                terms = tokenize(body) + tokenize(head) * 2 + tokenize(title_words)  # headings count extra
                chunks.append({"file": rel, "section": head, "text": body, "terms": Counter(terms),
                               "len": max(len(terms), 1)})
        heading_terms = Counter()  # topic words: used to judge one-word questions
        for c in chunks:
            heading_terms.update(set(tokenize(c["section"])) | set(tokenize(c["file"].rsplit(".", 1)[0])))
        postings = defaultdict(list)
        for i, c in enumerate(chunks):
            for term, tf in c["terms"].items():
                postings[term].append((i, tf))
        n = max(len(chunks), 1)
        idf = {t: math.log(1 + (n - len(p) + 0.5) / (len(p) + 0.5)) for t, p in postings.items()}
        avg = sum(c["len"] for c in chunks) / n
        for c in chunks:
            del c["terms"]  # not needed after indexing
        with self.lock:
            self.chunks, self.postings, self.idf, self.avg_len = chunks, dict(postings), idf, avg
            self.topics = {t for t, n in heading_terms.items() if n >= 3}
            self.files, self.errors = files, errors
            self.built_at, self.build_seconds = time.time(), round(time.time() - started, 1)
            self.building = False

    def build_in_background(self):
        threading.Thread(target=self.build, daemon=True).start()

    # ---------------------------------------------------------------- search
    def search(self, query, k=5, max_chars=7000, per_file=2):
        q = tokenize(query)
        if len(set(q)) == 1 and (q[0] not in self.topics or q[0] in GENERIC):
            return []  # one generic word ("better", "example") isn't enough to search on
        weights = {t: 2.0 for t in q}  # words actually said
        for t in q:
            for syn in SYNONYMS.get(t, []):
                weights.setdefault(syn, 0.6)  # related words help, but count once and less
        with self.lock:
            postings, idf, chunks, avg = self.postings, self.idf, self.chunks, self.avg_len
        if not chunks or not weights:
            return []
        k1, b = 1.4, 0.75
        scores = defaultdict(float)
        matched = defaultdict(set)
        for term, w in weights.items():
            for i, tf in postings.get(term, ()):
                L = chunks[i]["len"]
                scores[i] += w * idf[term] * tf * (k1 + 1) / (tf + k1 * (1 - b + b * L / avg))
                matched[i].add(term)
        if not scores:
            return []
        ranked = sorted(scores.items(), key=lambda x: -x[1])
        top = ranked[0][1]
        distinct_q = len(set(q))
        results, per, used, seen, labels = [], Counter(), 0, [], set()
        for i, s in ranked:
            if s < top * 0.45 or len(results) >= k:
                break
            # a match on just one common word is noise
            if distinct_q >= 2 and len(matched[i] & set(q)) < 2 and len(matched[i]) < 2:
                continue
            c = chunks[i]
            if per[c["file"]] >= per_file or used + len(c["text"]) > max_chars:
                continue
            words, label = set(tokenize(c["text"])), section_label(c).lower()
            if label in labels or any(len(words & w) / max(len(words | w), 1) > 0.6 for w in seen):
                continue  # same text in another file (e.g. a .md and .docx copy)
            seen.append(words)
            labels.add(label)
            per[c["file"]] += 1
            used += len(c["text"])
            results.append({"file": c["file"], "section": c["section"], "label": section_label(c),
                            "text": c["text"], "score": round(s, 2)})
        if results and results[0]["score"] < 4.0:  # weak best match: better to use none
            return []
        return results

    def status(self):
        with self.lock:
            return {"folder": str(self.root), "files": self.files, "sections": len(self.chunks),
                    "built_at": self.built_at, "build_seconds": self.build_seconds,
                    "building": self.building, "errors": len(self.errors)}
