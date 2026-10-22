"""Interview Copilot - local server.

Serves the web page and streams Claude's suggested answers.
Your API key stays here on the server and never reaches the browser.

Run:  python server.py     then open http://localhost:8765

For phone access, set ACCESS_CODE so only people with the code can use your API
credit (start-mobile.bat does this for you).
"""

import hmac
import io
import json
import os
import queue
import sys
import threading
import time
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlsplit

import anthropic

PORT = int(os.environ.get("PORT", "8765"))
HOST = os.environ.get("HOST", "127.0.0.1")  # 0.0.0.0 = reachable from other devices on your Wi-Fi
ACCESS_CODE = os.environ.get("ACCESS_CODE", "").strip()
MODEL = os.environ.get("COPILOT_MODEL", "claude-opus-5-5")
EFFORT = os.environ.get("COPILOT_EFFORT", "low")  # low = fastest live answers
STORY_EFFORT = os.environ.get("COPILOT_STORY_EFFORT", "medium")  # story bank is built once, so take more care
MAX_UPLOAD = 5 * 1024 * 1024
HERE = Path(__file__).parent
# Folder of resumes the page can load from. Default: the "Resume" folder next to this app's folder.
RESUME_DIR = Path(os.environ.get("RESUME_DIR", HERE.parent / "Resume"))
RESUME_TYPES = (".docx", ".pdf", ".md", ".txt")
# Resume, target role and story bank, shared by every device (kept out of git).
PROFILE_FILE = HERE / "data" / "profile.json"
PROFILE_FIELDS = ("profile", "role", "stories")

ANSWER_PROMPT = """You help a job candidate during a live interview. You receive a speech-to-text \
transcript of what the interviewer just said. The candidate's resume, target role and STAR story \
bank follow below. Latency-sensitive; begin your visible answer immediately.

Everything you write is something the candidate will say out loud: first person, natural spoken \
English, confident and honest. The transcript may contain recognition errors; infer the intended question.

## Behavioral, leadership and STAR questions
These include "tell me about a time...", "give me an example...", "describe a situation...", \
"how do you handle...", "what would your team say...", and anything about leadership, conflict, \
failure, mistakes, pressure, deadlines, ambiguity, influence, mentoring, disagreement, ownership, \
change or achievements.

Answer with ONE real story from the story bank, or straight from the resume if the bank is empty. \
Pick the story that best fits what the question is really testing. Avoid a story already used \
earlier in this interview when another one fits. Use this exact format:

**Story:** <story title> (<company>)
**Situation:** 1-2 sentences of context.
**Task:** 1 sentence: what you were responsible for.
**Action:**
- 2-4 bullets of what *you* did, starting with "I". For leadership questions show how you led \
people: set direction, aligned stakeholders, delegated, coached, made the call.
**Result:** the outcome, using the real numbers from the resume or story bank.
**Lesson:** one short sentence (only for leadership, failure, mistake or conflict questions).

Facts rule: use only companies, roles, projects, tools and numbers that appear in the resume or \
story bank. Never invent them. If a detail is missing, leave a visible placeholder such as \
[fill in: who disagreed] so the candidate can supply it. If no story truly fits, start with \
**Closest story:** instead of **Story:** and use the nearest one. Aim for about 180 words.

## All other questions
- Start with a one-sentence direct answer, then 2-4 short bullets, grounded in the resume where relevant.
- Technical questions: give the correct answer clearly; include a tiny code snippet only if it helps.
- "Tell me about yourself": a 4-5 sentence pitch from the resume, ending with why this target role.
- If the speech is not a question (small talk, instructions), reply with one short line suggesting how to respond.
- Keep it under about 150 words unless the question truly needs more.

Use simple Markdown (bold, bullets, `code`). No preamble like "Here's an answer"."""

STORY_PROMPT = """You turn a candidate's resume into a STAR story bank for behavioral and \
leadership interview questions.

Write 8 to 12 stories. Together they should cover as many of these themes as the resume \
supports: leadership, leading without authority, conflict, failure or mistake, tight deadline \
or pressure, ambiguity, mentoring or developing others, stakeholder or customer management, \
process improvement or automation, disagreeing with a manager, biggest achievement, learning \
something new fast. Prefer recent roles and stories with measurable results. If a target role \
is given, favour stories that matter most for it.

Format every story exactly like this, in first person:

### <number>. <Short title> - <Company>, <Role>, <Years>
**Themes:** <2-4 themes from the list>
- **Situation:** 1-2 sentences.
- **Task:** 1 sentence.
- **Action:** 2-3 sentences of what I did.
- **Result:** the outcome with numbers from the resume.

Facts rule: use only facts in the resume. Resumes rarely say who disagreed, what went wrong or \
how a conflict was resolved; where a story needs a detail like that, write a placeholder such as \
[fill in: what the disagreement was about] instead of inventing it.

After the stories, add a line "**Gaps:**" listing themes the resume gives no story for, so the \
candidate can prepare one. Output only the story bank."""


def clean_api_key():
    """Strip spaces and quotes that often sneak in when a key is pasted."""
    raw = os.environ.get("ANTHROPIC_API_KEY", "")
    key = raw.strip().strip('"').strip("'").strip()
    if key != raw:
        os.environ["ANTHROPIC_API_KEY"] = key
    return key, key != raw


def check_api_key():
    """Free call that tells us if the key works. Returns 'ok', 'missing', 'rejected' or 'unknown'."""
    if not have_key():
        return "missing"
    try:
        get_client().models.list(limit=1)
        return "ok"
    except anthropic.AuthenticationError:
        return "rejected"
    except anthropic.APIError:
        return "unknown"  # offline or temporary problem; answers may still work later


def have_key():
    return bool(os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN"))


KEY_STATUS = "unknown"
client = None


def get_client():
    global client
    if client is None:
        client = anthropic.Anthropic()
    return client


def extract_resume_text(filename, data):
    """Plain text from a .docx, .pdf, .txt or .md file."""
    name = filename.lower()
    if name.endswith(".docx"):
        import docx
        doc = docx.Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many resume templates put content in tables
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    if cell.text.strip() and cell.text not in cells:
                        cells.append(cell.text)
                lines.append(" | ".join(cells))
        text = "\n".join(lines)
    elif name.endswith(".pdf"):
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError("Please upload a .docx, .pdf or .txt file.")
    # tidy blank lines
    out, blank = [], False
    for line in text.splitlines():
        line = line.rstrip()
        if not line.strip():
            if not blank:
                out.append("")
            blank = True
        else:
            out.append(line)
            blank = False
    text = "\n".join(out).strip()
    if not text:
        raise ValueError("No text found in that file. If it's a scanned PDF, upload the .docx instead.")
    return text


def list_resumes():
    """Resume files in RESUME_DIR (skips Word's ~$ lock files)."""
    if not RESUME_DIR.is_dir():
        return []
    return sorted(p.name for p in RESUME_DIR.iterdir()
                  if p.is_file() and p.suffix.lower() in RESUME_TYPES and not p.name.startswith("~$"))


def candidate_context(profile, role, stories):
    parts = []
    if profile.strip():
        parts.append(f"<resume>\n{profile.strip()}\n</resume>")
    if role.strip():
        parts.append(f"<target_role>\n{role.strip()}\n</target_role>")
    if stories.strip():
        parts.append(f"<story_bank>\n{stories.strip()}\n</story_bank>")
    return "\n\n".join(parts) or "<resume>\n(The candidate has not added a resume yet.)\n</resume>"


def answer_request(question, profile, history):
    """System prompt + candidate context (cached between questions), then the live question."""
    system = [
        {"type": "text", "text": ANSWER_PROMPT},
        {
            "type": "text",
            "text": candidate_context(profile.get("profile", ""), profile.get("role", ""),
                                      profile.get("stories", "")),
            "cache_control": {"type": "ephemeral"},  # same for every question -> cheaper, faster
        },
    ]
    context = []
    if history:
        lines = [f"Q: {h['q']}\nA: {h['a']}" for h in history[-4:]]
        context.append("<earlier_in_this_interview>\n" + "\n\n".join(lines) + "\n</earlier_in_this_interview>")
    context.append(f"<interviewer_said>\n{question.strip()}\n</interviewer_said>")
    return system, [{"role": "user", "content": "\n\n".join(context)}]


def load_profile():
    try:
        data = json.loads(PROFILE_FILE.read_text(encoding="utf-8"))
        return {k: str(data.get(k, "")) for k in PROFILE_FIELDS}
    except (OSError, ValueError):
        return {k: "" for k in PROFILE_FIELDS}


# ---------------------------------------------------------------- usage metrics

# Estimated USD per million tokens: input, output, cache read, cache write (5-minute cache).
# Check https://www.anthropic.com/pricing if prices change.
PRICES = {
    "claude-opus-5-5": (4.00, 20.00, 0.20, 5.00),
    "claude-sonnet-5-5": (2.00, 10.00, 0.20, 2.50),
    "claude-haiku-5-5": (0.10, 0.50, 0.01, 0.125),
    "claude-opus-5": (5.00, 25.00, 0.50, 6.25),
    "claude-opus-4-8": (5.00, 25.00, 0.50, 6.25),
    "claude-fable-5-1": (10.00, 50.00, 0.25, 12.50),
}
USAGE_FILE = HERE / "data" / "usage.jsonl"
SETTINGS_FILE = HERE / "data" / "settings.json"


def price_for(model):
    if model in PRICES:
        return PRICES[model]
    for name in sorted(PRICES, key=len, reverse=True):
        if model.startswith(name):
            return PRICES[name]
    return PRICES["claude-opus-5-5"]


class Metrics:
    """Token use, cost, speed and rate limits for every call, saved in data/usage.jsonl."""

    def __init__(self):
        self.lock = threading.Lock()
        self.records = []
        self.limits = {}      # latest rate-limit headers from the API
        self.limits_at = None
        try:
            for line in USAGE_FILE.read_text(encoding="utf-8").splitlines():
                if line.strip():
                    self.records.append(json.loads(line))
        except (OSError, ValueError):
            pass
        try:
            self.settings = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            self.settings = {}
        if "budget" not in self.settings:
            self.settings["budget"] = float(os.environ.get("COPILOT_BUDGET_USD", "0") or 0)

    def set_budget(self, budget):
        with self.lock:
            self.settings["budget"] = max(0.0, float(budget))
            SETTINGS_FILE.parent.mkdir(exist_ok=True)
            SETTINGS_FILE.write_text(json.dumps(self.settings), encoding="utf-8")

    def update_limits(self, headers):
        found = {}
        for k, v in (headers or {}).items():
            k = k.lower()
            if k.startswith("anthropic-ratelimit-"):
                found[k[len("anthropic-ratelimit-"):]] = v
        if found:
            with self.lock:
                self.limits, self.limits_at = found, time.time()

    def record(self, kind, label, model, usage, ttft, total, status):
        u = {
            "input": getattr(usage, "input_tokens", 0) or 0,
            "output": getattr(usage, "output_tokens", 0) or 0,
            "cache_read": getattr(usage, "cache_read_input_tokens", 0) or 0,
            "cache_write": getattr(usage, "cache_creation_input_tokens", 0) or 0,
        }
        p = price_for(model)
        cost = (u["input"] * p[0] + u["output"] * p[1] + u["cache_read"] * p[2] + u["cache_write"] * p[3]) / 1e6
        rec = {"ts": round(time.time(), 1), "kind": kind, "label": label[:120], "model": model, **u,
               "cost": round(cost, 6), "saved": round(u["cache_read"] * (p[0] - p[2]) / 1e6, 6),
               "ttft": round(ttft, 2) if ttft is not None else None, "total": round(total, 2), "status": status}
        with self.lock:
            self.records.append(rec)
            USAGE_FILE.parent.mkdir(exist_ok=True)
            with USAGE_FILE.open("a", encoding="utf-8") as f:
                f.write(json.dumps(rec) + "\n")
        return rec

    def stats(self, interview_start):
        now = time.localtime()
        today = time.mktime((now.tm_year, now.tm_mon, now.tm_mday, 0, 0, 0, 0, 0, -1))
        month = time.mktime((now.tm_year, now.tm_mon, 1, 0, 0, 0, 0, 0, -1))
        with self.lock:
            records = list(self.records)
            limits, limits_at, budget = dict(self.limits), self.limits_at, self.settings.get("budget", 0)

        def summary(since):
            rs = [r for r in records if r["ts"] >= since]
            answers = [r for r in rs if r["kind"] == "answer"]
            tin = sum(r["input"] + r["cache_read"] + r["cache_write"] for r in rs)
            ttfts = [r["ttft"] for r in answers if r.get("ttft") is not None]
            totals = [r["total"] for r in answers if r["status"] == "ok"]
            return {
                "answers": len(answers),
                "stories": sum(1 for r in rs if r["kind"] == "stories"),
                "input": tin, "output": sum(r["output"] for r in rs),
                "cache_read": sum(r["cache_read"] for r in rs),
                "cost": round(sum(r["cost"] for r in rs), 4),
                "saved": round(sum(r.get("saved", 0) for r in rs), 4),
                "avg_ttft": round(sum(ttfts) / len(ttfts), 2) if ttfts else None,
                "avg_total": round(sum(totals) / len(totals), 2) if totals else None,
                "avg_cost": round(sum(r["cost"] for r in answers) / len(answers), 5) if answers else None,
            }

        periods = {"interview": summary(interview_start), "today": summary(today),
                   "month": summary(month), "all": summary(0)}
        spent = periods["month"]["cost"]
        recent = [{"ts": r["ts"], "label": r["label"], "in": r["input"] + r["cache_read"] + r["cache_write"],
                   "out": r["output"], "cost": r["cost"], "ttft": r.get("ttft"), "total": r["total"],
                   "kind": r["kind"], "status": r["status"]} for r in records[-20:]]
        return {"periods": periods, "budget": budget, "month_spent": spent,
                "budget_left": round(budget - spent, 4) if budget else None,
                "limits": limits, "limits_at": limits_at, "recent": recent, "model": MODEL}


METRICS = Metrics()


def stream_claude(system, messages, effort, max_tokens, emit, cancelled=lambda: False,
                  kind="answer", label=""):
    """Stream Claude's text to emit(); errors are emitted as readable text.
    Returns the usage record for this call (None if the call never ran)."""
    started, first, usage, served, status = time.time(), None, None, MODEL, "ok"
    try:
        with get_client().beta.messages.stream(
            model=MODEL,
            max_tokens=max_tokens,
            system=system,
            output_config={"effort": effort},
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
            messages=messages,
        ) as stream:
            try:
                METRICS.update_limits(stream.response.headers)
            except Exception:
                pass
            for text in stream.text_stream:
                if cancelled():
                    status = "cancelled"  # a newer question arrived; keep what was used so far
                    try:
                        usage = stream.current_message_snapshot.usage
                    except Exception:
                        pass
                    break  # leaving the block closes the connection to the API
                if first is None:
                    first = time.time()
                emit(text)
            if status == "ok":
                final = stream.get_final_message()
                usage, served = final.usage, final.model or MODEL
                if final.stop_reason == "refusal":
                    emit("\n\n_(Claude declined to answer this one.)_")
    except anthropic.AuthenticationError:
        emit("**Error:** Claude rejected your API key. Create a new key at console.anthropic.com "
             "(Settings → API Keys), copy all of it, then close the server window and start it again.")
    except anthropic.RateLimitError as e:
        METRICS.update_limits(getattr(e.response, "headers", {}))
        emit("**Error:** rate limited by the API. Wait a few seconds and try again.")
    except anthropic.APIStatusError as e:
        emit(f"**Error from the API ({e.status_code}):** {e.message}")
    except anthropic.APIConnectionError:
        emit("**Error:** could not reach the Claude API. Check your internet connection.")
    if usage is None:
        return None
    return METRICS.record(kind, label, served, usage,
                          (first - started) if first else None, time.time() - started, status)


class Session:
    """One shared interview: every open device sees the same questions, answers and resume."""

    def __init__(self):
        self.lock = threading.RLock()
        self.history = []        # finished [{id, q, a}], oldest first
        self.current = None      # {id, q, a, done}
        self.cancel = None       # stops the answer being written when a newer question arrives
        self.subscribers = set()
        self.profile = load_profile()
        self.started = time.time()  # start of this interview, for "This interview" stats

    def broadcast(self, event):
        with self.lock:
            subs = list(self.subscribers)
        for q in subs:
            q.put(event)

    def snapshot(self):
        with self.lock:
            snap = {"type": "snapshot", "history": [dict(h) for h in self.history[-30:]],
                    "current": dict(self.current) if self.current else None}
        snap["stats"] = METRICS.stats(self.started)
        return snap

    def ask(self, question, origin):
        with self.lock:
            if self.cancel:
                self.cancel.set()
            if self.current and self.current["a"].strip():
                c = self.current
                self.history = (self.history + [{"id": c["id"], "q": c["q"], "a": c["a"],
                                                 "usage": c.get("usage")}])[-50:]
            earlier = [dict(h) for h in self.history[-4:]]
            item = {"id": uuid.uuid4().hex[:12], "q": question, "a": "", "done": False}
            self.current, self.cancel = item, threading.Event()
            cancel, profile = self.cancel, dict(self.profile)
        self.broadcast({"type": "question", "id": item["id"], "q": question, "origin": origin})
        threading.Thread(target=self._write_answer, args=(item, cancel, profile, earlier), daemon=True).start()
        return item["id"]

    def _write_answer(self, item, cancel, profile, earlier):
        def emit(text):
            with self.lock:
                item["a"] += text
            self.broadcast({"type": "delta", "id": item["id"], "text": text})

        system, messages = answer_request(item["q"], profile, earlier)
        rec = stream_claude(system, messages, EFFORT, 4000, emit, cancel.is_set, kind="answer", label=item["q"])
        summary = None
        if rec:
            summary = {"in": rec["input"] + rec["cache_read"] + rec["cache_write"], "cached": rec["cache_read"],
                       "written": rec["cache_write"],
                       "out": rec["output"], "cost": rec["cost"], "ttft": rec["ttft"], "total": rec["total"]}
        with self.lock:
            item["done"] = True
            item["usage"] = summary
        self.broadcast({"type": "done", "id": item["id"], "usage": summary})
        if rec:
            self.broadcast_stats()

    def broadcast_stats(self):
        self.broadcast({"type": "stats", "stats": METRICS.stats(self.started)})

    def reset(self):
        with self.lock:
            if self.cancel:
                self.cancel.set()
            self.history, self.current, self.cancel = [], None, None
            self.started = time.time()
        self.broadcast({"type": "reset"})
        self.broadcast_stats()

    def save_profile(self, fields, origin):
        with self.lock:
            for k in PROFILE_FIELDS:
                if k in fields:
                    self.profile[k] = str(fields[k])
            PROFILE_FILE.parent.mkdir(exist_ok=True)
            tmp = PROFILE_FILE.with_suffix(".tmp")
            tmp.write_text(json.dumps(self.profile, ensure_ascii=False, indent=1), encoding="utf-8")
            tmp.replace(PROFILE_FILE)
            changed = {k: self.profile[k] for k in PROFILE_FIELDS if k in fields}
        self.broadcast({"type": "profile", "origin": origin, **changed})


SESSION = Session()


def story_request(data):
    content = f"<resume>\n{str(data.get('profile', '')).strip()}\n</resume>"
    role = str(data.get("role", "")).strip()
    if role:
        content += f"\n\n<target_role>\n{role}\n</target_role>"
    return STORY_PROMPT, [{"role": "user", "content": content}]


class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"  # needed for chunked streaming

    def log_message(self, fmt, *args):
        if "/events" not in (args[0] if args else ""):  # the live feed would flood the log
            sys.stderr.write("[server] " + (fmt % args) + "\n")

    def _authorized(self):
        """True when no access code is configured, or the request carries the right one."""
        if not ACCESS_CODE:
            return True
        sent = self.headers.get("X-Access-Code", "").strip()
        if not sent:  # the live feed can't send headers, so it passes ?code=
            sent = (parse_qs(urlsplit(self.path).query).get("code") or [""])[0].strip()
        if hmac.compare_digest(sent.encode(), ACCESS_CODE.encode()):
            return True
        time.sleep(1)  # slows down anyone guessing codes
        self.close_connection = True  # any unread request body is discarded with the connection
        self._json(401, {"error": "Wrong or missing access code."})
        return False

    def do_GET(self):
        url = urlsplit(self.path)
        if url.path in ("/", "/index.html"):
            body = (HERE / "index.html").read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-cache")
            self.end_headers()
            self.wfile.write(body)
            return
        if not self._authorized():
            return
        if url.path == "/status":
            self._json(200, {"model": MODEL, "effort": EFFORT, "key_configured": have_key(),
                             "key_status": KEY_STATUS})
        elif url.path == "/events":
            self._events()
        elif url.path == "/profile":
            with SESSION.lock:
                self._json(200, dict(SESSION.profile))
        elif url.path == "/stats":
            self._json(200, METRICS.stats(SESSION.started))
        elif url.path == "/resumes":
            self._json(200, {"folder": str(RESUME_DIR), "files": list_resumes()})
        elif url.path == "/resume":
            name = unquote((parse_qs(url.query).get("name") or [""])[0])
            if name not in list_resumes():  # only files from the listing, never arbitrary paths
                self._json(404, {"error": "Resume not found."})
                return
            try:
                text = extract_resume_text(name, (RESUME_DIR / name).read_bytes())
            except Exception as e:
                self._json(400, {"error": f"Couldn't read {name}: {e}"})
                return
            self._json(200, {"text": text})
        else:
            self._json(404, {"error": "Not found"})

    def _events(self):
        """Live feed (Server-Sent Events): every device gets each question and answer as it's written."""
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream; charset=utf-8")
        self.send_header("Cache-Control", "no-cache")
        self.send_header("Transfer-Encoding", "chunked")
        self.end_headers()
        q = queue.Queue()
        with SESSION.lock:
            SESSION.subscribers.add(q)
            first = SESSION.snapshot()
        try:
            self._chunk("retry: 2000\n\n")
            self._chunk("data: " + json.dumps(first) + "\n\n")
            while True:
                try:
                    event = q.get(timeout=15)
                    self._chunk("data: " + json.dumps(event) + "\n\n")
                except queue.Empty:
                    self._chunk(": ping\n\n")  # keeps phone/tunnel connections open
        except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError, OSError):
            pass
        finally:
            with SESSION.lock:
                SESSION.subscribers.discard(q)
            self.close_connection = True

    def do_POST(self):
        path = urlsplit(self.path).path
        if path not in ("/answer", "/stories", "/extract", "/profile", "/reset", "/settings"):
            self._json(404, {"error": "Not found"})
            return
        if not self._authorized():
            return
        length = int(self.headers.get("Content-Length", "0") or 0)
        if length > MAX_UPLOAD:
            self.close_connection = True
            self._json(413, {"error": "That file is too big (5 MB max)."})
            return
        raw = self.rfile.read(length)

        if path == "/extract":
            try:
                text = extract_resume_text(self.headers.get("X-Filename", ""), raw)
            except ValueError as e:
                self._json(400, {"error": str(e)})
                return
            except Exception:
                self._json(400, {"error": "Couldn't read that file. Try saving it as .docx or .pdf again."})
                return
            self._json(200, {"text": text})
            return

        try:
            data = json.loads(raw or b"{}")
        except (ValueError, json.JSONDecodeError):
            self._json(400, {"error": "Request body must be JSON."})
            return
        origin = str(data.get("origin", ""))

        if path == "/profile":
            SESSION.save_profile({k: data[k] for k in PROFILE_FIELDS if k in data}, origin)
            self._json(200, {"ok": True})
            return
        if path == "/reset":
            SESSION.reset()
            self._json(200, {"ok": True})
            return
        if path == "/settings":
            try:
                METRICS.set_budget(data.get("budget", 0))
            except (TypeError, ValueError):
                self._json(400, {"error": "Budget must be a number, e.g. 10"})
                return
            SESSION.broadcast_stats()
            self._json(200, {"ok": True})
            return

        if not have_key():
            self._json(500, {"error": "No API key set. Set ANTHROPIC_API_KEY and restart the server."})
            return

        if path == "/answer":
            question = str(data.get("question", "")).strip()
            if not question:
                self._json(400, {"error": "No question was sent."})
                return
            self._json(202, {"id": SESSION.ask(question, origin)})  # the answer arrives on /events
            return

        # /stories: streamed back to the device that asked
        if not str(data.get("profile", "")).strip():
            self._json(400, {"error": "Add your resume first."})
            return
        system, messages = story_request(data)
        self.send_response(200)
        # event-stream type stops tunnels/proxies from buffering; the page reads it as plain text
        self.send_header("Content-Type", "text/event-stream; charset=utf-8")
        self.send_header("Cache-Control", "no-cache")
        self.send_header("Transfer-Encoding", "chunked")
        self.end_headers()
        stopped = threading.Event()

        def send(text):
            try:
                self._chunk(text)
            except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError):
                stopped.set()  # the page pressed Stop

        rec = stream_claude(system, messages, STORY_EFFORT, 16000, send, stopped.is_set,
                            kind="stories", label="Story bank")
        if rec:
            SESSION.broadcast_stats()
        if not stopped.is_set():
            try:
                self.wfile.write(b"0\r\n\r\n")
            except OSError:
                pass

    def _chunk(self, text):
        b = text.encode("utf-8")
        if b:
            self.wfile.write(f"{len(b):X}\r\n".encode() + b + b"\r\n")
            self.wfile.flush()

    def _json(self, code, obj):
        body = json.dumps(obj).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


if __name__ == "__main__":
    key, cleaned = clean_api_key()
    if cleaned:
        print("Note: removed extra spaces/quotes from the pasted API key.")
    print("Checking your API key...")
    KEY_STATUS = check_api_key()
    if KEY_STATUS == "ok":
        print("API key works.")
    elif KEY_STATUS == "missing":
        print("Warning: ANTHROPIC_API_KEY is not set - answers will fail until you set it.")
    elif KEY_STATUS == "rejected":
        looks = "starts with sk-ant-" if key.startswith("sk-ant-") else "does NOT start with sk-ant- (likely the wrong text was pasted)"
        print("=" * 64)
        print("PROBLEM: Claude rejected this API key.")
        print(f"  The key you entered is {len(key)} characters long and {looks}.")
        print("  Real keys start with sk-ant- and are usually 100+ characters.")
        print("  Fix: create a new key at console.anthropic.com > Settings > API Keys,")
        print("  copy the whole key, close this window and start again.")
        print("=" * 64)
    else:
        print("Couldn't check the API key right now (no internet?). Answers will try anyway.")
    if HOST != "127.0.0.1" and not ACCESS_CODE:
        print("Warning: the server is reachable from other devices but ACCESS_CODE is not set.")
    print(f"Interview Copilot running at http://localhost:{PORT}  (model: {MODEL}, effort: {EFFORT})")
    if ACCESS_CODE:
        print(f"Access code: {ACCESS_CODE}")
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
